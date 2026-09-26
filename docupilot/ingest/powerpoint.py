"""PPTX text, tables, native chart values, pictures and notes with exact slide citations."""

from collections.abc import Callable
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from docupilot.ingest.text import table_markdown
from docupilot.models import Section

PictureParser = Callable[[bytes, int, str], list[Section]]


def parse_pptx(path: Path, file_name: str, *, picture_parser: PictureParser | None = None) -> list[Section]:
    presentation = Presentation(path)
    sections: list[Section] = []
    for number, slide in enumerate(presentation.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        heading = f"Slide {number}" + (f" > {title}" if title else "")

        def emit(body: str, *, kind="text", suffix="") -> None:
            if body.strip():
                sections.append(Section(body.strip(), file_name, page=number,
                                        heading_path=heading + suffix, kind=kind))

        def walk(shapes) -> None:
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(shape.shapes)
                elif shape.has_table:
                    emit(table_markdown([[cell.text for cell in row.cells] for row in shape.table.rows]), kind="table")
                elif shape.has_chart:
                    chart = shape.chart
                    chart_title = chart.chart_title.text_frame.text if chart.has_title else "Chart"
                    labels = []
                    for name in ("category_axis", "value_axis"):
                        try:
                            axis = getattr(chart, name)
                        except ValueError:  # pie and other charts have no category/value axes
                            continue
                        if axis.has_title:
                            labels.append(f"{name.replace('_', ' ')}: {axis.axis_title.text_frame.text}")
                    emit("\n".join(labels), suffix=f" > {chart_title}")
                    for plot in chart.plots:
                        categories = [str(c.label) for c in getattr(plot, "categories", [])]
                        for series in plot.series:
                            xs = list(getattr(series, "x_values", categories))
                            rows = [["Category / X", str(series.name)]]
                            rows += [[str(xs[i]) if i < len(xs) else str(i + 1), str(value)]
                                     for i, value in enumerate(series.values)]
                            emit(table_markdown(rows), kind="table", suffix=f" > {chart_title}")
                elif picture_parser and hasattr(shape, "image"):
                    sections.extend(picture_parser(shape.image.blob, number, heading + " > Picture"))
                elif shape.has_text_frame:
                    emit(shape.text)

        walk(slide.shapes)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None:
                emit(notes.text, suffix=" > Speaker notes")
    return sections
