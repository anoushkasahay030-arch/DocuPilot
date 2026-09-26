"""Synthetic format and workspace regressions; no model downloads or live inference."""

from io import BytesIO
from pathlib import Path
import re
import tomllib
from unittest.mock import Mock

import httpx
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
import pymupdf
import pytest

from docupilot.config import Settings
from docupilot.ingest import code, html, image, powerpoint, web
from docupilot.ingest.chunker import chunk_sections
from docupilot.ingest.formats import CODE_NAMES, CODE_TYPES, IMAGE_TYPES, SUPPORTED
from docupilot.models import Citation
from docupilot.workspace import Workspace


@pytest.fixture
def settings(tmp_path):
    return Settings(_env_file=None, data_dir=tmp_path, ollama_vision_model="test-vision")


@pytest.fixture
def workspace(settings):
    store = Mock()
    ws = Workspace("formats", store, settings)
    yield ws
    ws.close()


def png_bytes(color="white"):
    output = BytesIO()
    Image.new("RGB", (100, 50), color).save(output, format="PNG")
    return output.getvalue()


@pytest.fixture
def deck(tmp_path):
    path = tmp_path / "roadmap.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Launch plan"
    slide.placeholders[1].text = "Release in October."
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = "Owner: Ada"
    table = slide.shapes.add_table(2, 2, 0, 0, Inches(3), Inches(1)).table
    for cell, value in zip([table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)],
                           ["Region", "Revenue", "West", "42"]):
        cell.text = value
    slide.notes_slide.notes_text_frame.text = "Review on Friday."
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("Sales", [10, 20])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(4), Inches(3), data).chart
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = "Revenue ($M)"
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    slide2.shapes.add_picture(BytesIO(png_bytes()), 0, 0)
    prs.slides.add_slide(prs.slide_layouts[6])  # empty slide must still count
    prs.save(path)
    return path


def test_powerpoint_text_tables_charts_notes_and_picture_locations(deck):
    calls = []
    def picture(data, page, heading):
        calls.append((data, page, heading))
        return []
    sections = powerpoint.parse_pptx(deck, deck.name, picture_parser=picture)
    assert all(s.file_name == deck.name and s.page == 1 for s in sections)
    assert any("Owner: Ada" in s.text for s in sections)
    assert any("Review on Friday" in s.text and s.heading_path.endswith("Speaker notes") for s in sections)
    assert any("|West|42|" in s.text and s.kind == "table" for s in sections)
    assert any("|Q2|20.0|" in s.text and "Chart" in s.heading_path for s in sections)
    assert any("Revenue ($M)" in s.text for s in sections)
    assert calls[0][1:] == (2, "Slide 2 > Picture")
    assert Citation(1, deck.name, "text", page=2).label() == "roadmap.pptx, slide 2"


def test_powerpoint_template_picture_is_extracted(tmp_path):
    path = tmp_path / "template.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[8])  # picture-with-caption template
    slide.placeholders[1].insert_picture(BytesIO(png_bytes()))
    prs.save(path)
    parse_picture = Mock(return_value=[])
    powerpoint.parse_pptx(path, path.name, picture_parser=parse_picture)
    parse_picture.assert_called_once()
    assert parse_picture.call_args.args[1] == 1


def test_workspace_pptx_preserves_partial_content_and_reports_vision_failure(workspace, deck, monkeypatch):
    monkeypatch.setattr(image, "describe_image", Mock(side_effect=RuntimeError("ollama pull test-vision")))
    info = workspace.ingest(deck)
    assert info.pages == 3 and info.chunks > 0
    assert "ollama pull test-vision" in info.warnings[0]
    chunks = workspace.store.upsert.call_args.args[0]
    assert all(c.session_id == "formats" and c.file_id == info.file_id for c in chunks)


def test_html_headings_tables_code_and_hidden_content(tmp_path):
    path = tmp_path / "page.html"
    path.write_text('''<html><head><title>Manual</title><script>secret()</script></head>
      <body><nav>MENU</nav><h1>Guide</h1><p>Read <b>carefully</b>.</p><h2>Costs</h2>
      <table><tr><th>Item</th><th>Price</th></tr><tr><td>Widget</td><td>42</td></tr></table>
      <pre>def run():\n    return 42</pre><img src="https://example.test/x" alt="Revenue chart">
      <div hidden>HIDDEN</div><div style="display: none">INVISIBLE</div>
      <script>BAD_SCRIPT</script><!-- COMMENT --></body></html>''')
    sections = html.parse_html(path, path.name)
    body = "\n".join(s.text for s in sections)
    assert "Read carefully." in body and "Revenue chart" in body
    assert not any(token in body for token in ["secret", "MENU", "HIDDEN", "INVISIBLE", "BAD_SCRIPT", "COMMENT"])
    table = next(s for s in sections if s.kind == "table")
    assert table.heading_path == "Guide > Costs" and "|Widget|42|" in table.text
    source = next(s for s in sections if s.kind == "code")
    assert "    return 42" in source.text and source.heading_path.endswith("lines 1–2")


def test_code_keeps_whitespace_and_line_citations_without_execution(tmp_path):
    path = tmp_path / "source.py"
    path.write_text("# comment, not heading\ndef run():\n    return 42\n\nraise RuntimeError('never execute')\n")
    sections = code.parse_code(path, path.name, target_tokens=18)
    chunks = chunk_sections(sections, file_id="f", session_id="s", target_tokens=18)
    assert all(c.kind == "code" and c.heading_path.startswith("lines ") for c in chunks)
    assert any("    return 42" in c.text for c in chunks)
    assert sum("raise RuntimeError" in c.text for c in chunks) == 1
    assert all(c.text.startswith("```\n") and c.text.endswith("\n```") for c in chunks)


def test_minified_code_is_bounded_and_not_lost():
    source = "x=1;" * 500
    sections = code.code_sections(source, "app.js", target_tokens=30)
    assert all(len(s.text) <= 30 * 4 + 8 for s in sections)
    assert "".join(s.text[4:-4] for s in sections) == source
    assert all(s.heading_path == "lines 1–1" for s in sections)


def test_utf16_code_and_binary_rejection(tmp_path):
    path = tmp_path / "test.py"
    path.write_bytes("name = 'café'\n".encode("utf-16"))
    assert "café" in code.parse_code(path, path.name)[0].text
    path.write_bytes(b"binary\x00data")
    with pytest.raises(ValueError, match="binary"):
        code.parse_code(path, path.name)


@pytest.mark.parametrize("name", ["test" + suffix.upper() for suffix in sorted(CODE_TYPES)] + sorted(CODE_NAMES))
def test_code_formats_reach_workspace(workspace, tmp_path, name):
    path = tmp_path / name
    path.write_text("value = 42\n")
    info = workspace.ingest(path)
    assert info.kind == "document" and info.chunks == 1
    assert workspace.store.upsert.call_args.args[0][0].kind == "code"


@pytest.mark.parametrize("suffix", sorted(IMAGE_TYPES))
def test_image_formats_normalized_and_transcribed(workspace, tmp_path, monkeypatch, suffix):
    path = tmp_path / ("scan" + suffix)
    Image.new("RGB", (100, 50), "white").save(path)
    describe = Mock(return_value=image.ImageEvidence(transcription="Revenue: 42", description="An ascending bar chart."))
    monkeypatch.setattr(image, "describe_image", describe)
    info = workspace.ingest(path)
    assert info.chunks == 2
    chunks = workspace.store.upsert.call_args.args[0]
    assert {c.text for c in chunks} == {"Revenue: 42", "An ascending bar chart."}
    assert all("model-generated" in c.heading_path for c in chunks)
    assert Image.open(BytesIO(describe.call_args.args[0])).format == "PNG"


def test_multiframe_images_keep_frame_locations_and_enforce_limit(settings, monkeypatch):
    output = BytesIO()
    Image.new("RGB", (10, 10), "white").save(output, format="TIFF", save_all=True,
                                            append_images=[Image.new("RGB", (10, 10), "black")])
    describe = Mock(return_value=image.ImageEvidence(transcription="text", description=""))
    monkeypatch.setattr(image, "describe_image", describe)
    sections = image.parse_image_bytes(output.getvalue(), "scan.tiff", settings)
    assert [s.heading_path.split(" > ")[0] for s in sections] == ["frame 1", "frame 2"]
    settings.image_max_frames = 1
    with pytest.raises(ValueError, match="IMAGE_MAX_FRAMES"):
        image.parse_image_bytes(output.getvalue(), "scan.tiff", settings)
    assert describe.call_count == 2


def test_image_resizing_and_transparency(settings, monkeypatch):
    settings.image_max_side = 256
    output = BytesIO()
    Image.new("RGBA", (1024, 512), (0, 0, 0, 0)).save(output, format="PNG")
    describe = Mock(return_value=image.ImageEvidence(transcription="", description="Blank image"))
    monkeypatch.setattr(image, "describe_image", describe)
    image.parse_image_bytes(output.getvalue(), "scan.png", settings)
    normalized = Image.open(BytesIO(describe.call_args.args[0]))
    assert normalized.size == (256, 128)
    assert normalized.getpixel((0, 0)) == (255, 255, 255)


def test_bad_pptx_picture_does_not_skip_later_valid_pictures(workspace, deck, monkeypatch):
    prs = Presentation(deck)
    prs.slides[1].shapes.add_picture(BytesIO(png_bytes("black")), 0, 0)
    prs.save(deck)
    parser = image.parse_image_bytes
    calls = []
    def read(data, *args, **kwargs):
        calls.append(data)
        if len(calls) == 1:
            raise OSError("Unsupported image encoding")
        return parser(data, *args, **kwargs)
    monkeypatch.setattr(image, "parse_image_bytes", read)
    monkeypatch.setattr(image, "describe_image", Mock(return_value=image.ImageEvidence(
        transcription="Visible in the second picture", description="")))
    info = workspace.ingest(deck)
    assert len(info.warnings) == 1
    assert any("second picture" in c.text for c in workspace.store.upsert.call_args.args[0])


def test_pure_scanned_pdf_reports_missing_vision_and_is_not_registered(workspace, tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_image(page.rect, stream=png_bytes())
        doc.save(path)
    monkeypatch.setattr(image, "describe_image", Mock(side_effect=RuntimeError("ollama pull test-vision")))
    with pytest.raises(ValueError, match="ollama pull test-vision"):
        workspace.ingest(path)
    assert not workspace.files


def test_missing_vision_model_does_not_register_image(workspace, tmp_path, monkeypatch):
    path = tmp_path / "scan.png"
    path.write_bytes(png_bytes())
    monkeypatch.setattr(image, "describe_image", Mock(side_effect=RuntimeError("ollama pull test-vision")))
    with pytest.raises(RuntimeError, match="ollama pull"):
        workspace.ingest(path)
    assert not workspace.files
    workspace.store.upsert.assert_not_called()


def test_scanned_pdf_uses_vision_preserves_page_and_skips_blank(workspace, tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_image(page.rect, stream=png_bytes())
        doc.new_page()  # empty pages should not invoke vision
        doc.new_page().insert_text((72, 300), "Native text on page three.\nThe invoice was paid in October.")
        doc.save(path)
    describe = Mock(return_value=image.ImageEvidence(transcription="Scanned invoice: 42", description=""))
    monkeypatch.setattr(image, "describe_image", describe)
    info = workspace.ingest(path)
    chunks = workspace.store.upsert.call_args.args[0]
    assert info.pages == 3
    assert next(c for c in chunks if "invoice" in c.text).page == 1
    assert next(c for c in chunks if "Native text" in c.text).page == 3
    assert describe.call_count == 1


def test_web_import_redirects_and_keeps_origin(workspace, monkeypatch):
    requests = []
    def handle(request):
        requests.append(request)
        if request.url.path == "/old":
            return httpx.Response(302, headers={"location": "/guide"})
        return httpx.Response(200, headers={"content-type": "text/html; charset=utf-8"},
                              content=b"<h1>Guide</h1><p>The limit is 42.</p>")
    fetch = web.fetch_page
    monkeypatch.setattr(web, "fetch_page", lambda url, **kw: fetch(url, **kw, transport=httpx.MockTransport(handle)))
    info = workspace.ingest_url("https://example.test/old")
    assert info.source_url == "https://example.test/guide"
    assert Path(info.path).is_file()
    chunks = workspace.store.upsert.call_args.args[0]
    assert all(c.heading_path.startswith(info.source_url) for c in chunks)
    assert len(requests) == 2
    workspace.ingest_url("https://example.test/guide")
    assert len(workspace.files) == 1
    assert len(list(workspace.settings.uploads_dir.rglob("*.html"))) == 1


@pytest.mark.parametrize("url", ["file:///etc/passwd", "ftp://example.test/page", "https://u:p@example.test", ""])
def test_url_rejects_non_http_and_credentials(url):
    with pytest.raises(ValueError, match="HTTP|http"):
        web.fetch_page(url, max_bytes=100)


@pytest.mark.parametrize("response, message", [
    (httpx.Response(200, headers={"content-type": "text/html"}, content=b"x" * 101), "size limit"),
    (httpx.Response(200, headers={"content-type": "application/pdf"}, content=b"pdf"), "does not serve HTML"),
    (httpx.Response(302, headers={"location": "file:///tmp/file"}), "http"),
    (httpx.Response(302, headers={"location": "/loop"}), "too many redirects"),
])
def test_web_response_limits(response, message):
    with pytest.raises(ValueError, match=message):
        web.fetch_page("https://example.test", max_bytes=100, transport=httpx.MockTransport(lambda _: response))


def test_empty_file_is_not_registered(workspace, tmp_path):
    path = tmp_path / "empty.html"
    path.write_text("<html><script>ignored()</script></html>")
    with pytest.raises(ValueError, match="no extractable content"):
        workspace.ingest(path)
    assert not workspace.files


def test_upload_picker_covers_backend_extensions():
    config = tomllib.loads((Path(__file__).parents[1] / ".chainlit/config.toml").read_text())
    accepted = config["features"]["spontaneous_file_upload"]["accept"]
    # react-dropzone 14.2.3 discards the entire entry (including its extensions)
    # for application/*. Its native picker accepts only these wildcard types.
    for mime in accepted:
        assert mime in {"audio/*", "video/*", "image/*", "text/*"} or re.fullmatch(
            r"\w+/[-+.\w]+", mime
        ), f"The native file picker discards {mime!r}"
    assert SUPPORTED <= {extension for extensions in accepted.values() for extension in extensions}


@pytest.mark.parametrize("name, mime", [
    ("report.pdf", "application/pdf"),
    ("handbook.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ("sales.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ("sales.xlsm", "application/vnd.ms-excel.sheet.macroenabled.12"),
    ("sales.xlsm", "application/vnd.ms-excel.sheet.macroEnabled.12"),
    ("customers.csv", "application/vnd.ms-excel"),
    ("handbook.docx", "application/zip"),
    ("report.PDF", "application/octet-stream"),
    ("source.py", "text/x-python"), ("source.js", "text/javascript"),
    ("source.js", "application/javascript"), ("source.ts", "video/mp2t"),
    ("source.json", "application/json"), ("source.xml", "text/xml"),
    ("Dockerfile", "application/octet-stream"), ("Makefile", "text/x-makefile"),
    ("slide.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    ("page.xhtml", "application/xhtml+xml"), ("scan.png", "image/png"),
    ("source.php", "application/x-httpd-php"), ("source.sh", "application/x-sh"),
    ("source.sql", "application/x-sql"), ("source.rs", "application/rls-services+xml"),
])
def test_chainlit_accepts_browser_mime_variants(name, mime, monkeypatch):
    # Exercise the installed server's actual validator; matching the extension alone is insufficient.
    monkeypatch.delenv("DEBUG", raising=False)
    from chainlit.server import validate_file_mime_type
    from fastapi import UploadFile
    from starlette.datastructures import Headers
    file = UploadFile(BytesIO(b"test"), filename=name, headers=Headers({"content-type": mime}))
    validate_file_mime_type(file, None)
