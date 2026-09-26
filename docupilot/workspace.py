"""A chat session's files: parsing → chunking → indexing, plus the session's spreadsheet tables."""

import hashlib
import logging
import shutil
from collections.abc import Callable
from pathlib import Path
from uuid import uuid4

from pptx import Presentation

from docupilot.config import Settings, get_settings
from docupilot.index.store import VectorStore
from docupilot.ingest import code, html, image, pdf, powerpoint, text, web, word
from docupilot.ingest.chunker import chunk_sections
from docupilot.ingest.formats import CODE_NAMES, CODE_TYPES, DOC_TYPES, HTML_TYPES, IMAGE_TYPES, SHEET_TYPES, SUPPORTED
from docupilot.ingest.tabular import TableStore
from docupilot.models import FileInfo, Section

log = logging.getLogger(__name__)

Progress = Callable[[str], None]


class UnsupportedFile(ValueError):
    pass


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()[:16]


class Workspace:
    def __init__(self, session_id: str, store: VectorStore, settings: Settings | None = None):
        self.session_id = session_id
        self.store = store
        self.settings = settings or get_settings()
        self.tables = TableStore()
        self.files: dict[str, FileInfo] = {}
        self.web_urls: dict[str, str] = {}

    @property
    def doc_files(self) -> list[FileInfo]:
        return [f for f in self.files.values() if f.kind == "document"]

    @property
    def sheet_files(self) -> list[FileInfo]:
        return [f for f in self.files.values() if f.kind == "spreadsheet"]

    def file_by_name(self, name: str) -> FileInfo | None:
        return next((f for f in self.files.values() if f.name == name), None)

    def ingest(self, path: Path, file_name: str | None = None, progress: Progress | None = None,
               *, source_url: str | None = None) -> FileInfo:
        path = Path(path)
        file_name = file_name or path.name
        ext = Path(file_name).suffix.lower()
        say = progress or (lambda _msg: None)
        is_code = ext in CODE_TYPES or Path(file_name).name.lower() in CODE_NAMES
        if ext not in SUPPORTED and not is_code:
            raise UnsupportedFile(f"{file_name}: unsupported type {ext or '(none)'}. "
                                  f"Supported: {', '.join(sorted(SUPPORTED))}")
        size_mb = path.stat().st_size / 1e6
        if size_mb > self.settings.max_file_mb:
            raise UnsupportedFile(f"{file_name}: {size_mb:.0f} MB exceeds the {self.settings.max_file_mb} MB limit")
        file_id = file_hash(path)
        if file_id in self.files:
            say(f"{file_name}: already indexed")
            return self.files[file_id]
        if len(self.files) >= self.settings.max_files:
            raise UnsupportedFile(f"File limit reached ({self.settings.max_files}).")
        if self.file_by_name(file_name):
            # Same name, different content: make the citation label unambiguous.
            file_name = f"{Path(file_name).stem} ({file_id[:6]}){ext}"

        info = FileInfo(file_id=file_id, name=file_name, path=str(path),
                        kind="spreadsheet" if ext in SHEET_TYPES else "document", source_url=source_url)
        vision_unavailable = False

        def visual_parser(data: bytes, page: int, heading: str) -> list[Section]:
            nonlocal vision_unavailable
            # Preserve readable document content when the optional vision service is unavailable.
            if vision_unavailable:
                return []
            say(f"{file_name}: reading image on page/slide {page} with {self.settings.ollama_vision_model}")
            try:
                return image.parse_image_bytes(data, file_name, self.settings, page=page, heading_path=heading)
            except RuntimeError as error:
                vision_unavailable = True
                warning = f"Visual content was skipped from page/slide {page} onward: {error}"
            except (ValueError, OSError) as error:
                warning = f"Could not read a picture on page/slide {page}: {error}"
            info.warnings.append(warning)
            say(warning)
            return []

        sections: list[Section] = []
        if ext in SHEET_TYPES:
            tables = self.tables.add_file(path, file_name)
            info.tables = [t.name for t in tables]
            sections = [Section(text=t.schema_card, file_name=file_name, page=None,
                                heading_path=f"sheet {t.sheet}" if t.sheet else "", kind="table_schema")
                        for t in tables]
            say(f"{file_name}: loaded {len(tables)} table(s): "
                + ", ".join(f"{t.name} ({t.row_count} rows)" for t in tables))
        elif ext == ".pdf":
            for page_no, total, secs in pdf.iter_pdf_sections(path, file_name, use_ocr=self.settings.use_ocr,
                                                            visual_parser=visual_parser):
                sections += secs
                info.pages = total
                if page_no % 25 == 0 or page_no == total:
                    say(f"{file_name}: parsed page {page_no}/{total}")
        elif ext == ".docx":
            sections = word.parse_docx(path, file_name)
            pages = [s.page for s in sections if s.page]
            info.pages = max(pages) if pages else None
        elif ext == ".pptx":
            sections = powerpoint.parse_pptx(path, file_name, picture_parser=visual_parser)
            info.pages = len(Presentation(path).slides)
        elif ext in IMAGE_TYPES:
            say(f"{file_name}: extracting text and visual evidence with {self.settings.ollama_vision_model}")
            sections = image.parse_image(path, file_name, self.settings)
        elif ext in HTML_TYPES:
            sections = html.parse_html(path, file_name, target_tokens=self.settings.chunk_tokens)
        elif is_code:
            sections = code.parse_code(path, file_name, target_tokens=self.settings.chunk_tokens)
        else:
            sections = text.parse_text_file(path, file_name)

        if source_url:
            for section in sections:
                section.heading_path = " > ".join(filter(None, [source_url, section.heading_path]))

        s = self.settings
        chunks = chunk_sections(sections, file_id=file_id, session_id=self.session_id, target_tokens=s.chunk_tokens,
                                overlap_tokens=s.chunk_overlap_tokens, max_table_tokens=s.max_table_chunk_tokens)
        if not chunks:
            detail = " ".join(info.warnings)
            raise ValueError(f"{file_name}: no extractable content. {detail}".strip())
        self.store.upsert(chunks, on_progress=lambda done, total: say(f"{file_name}: indexed {done}/{total} chunks")
                          if total > 64 and (done == total or done % 256 == 0) else None)
        info.chunks = len(chunks)
        self.files[file_id] = info
        return info

    def ingest_url(self, url: str, progress: Progress | None = None) -> FileInfo:
        key = web.url_key(url)
        existing = self.files.get(self.web_urls.get(key, ""))
        if existing:
            if progress:
                progress(f"{existing.name}: already loaded in this chat")
            return existing
        if len(self.files) >= self.settings.max_files:
            raise UnsupportedFile(f"File limit reached ({self.settings.max_files}).")
        page = web.fetch_page(url, max_bytes=self.settings.max_file_mb * 1_000_000)
        directory = self.settings.uploads_dir / self.session_id / uuid4().hex
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / page.file_name
        try:
            path.write_bytes(page.content)
            info = self.ingest(path, page.file_name, progress, source_url=page.url)
        except Exception:
            shutil.rmtree(directory, ignore_errors=True)
            raise
        if info.path != str(path):  # duplicate content already belongs to this chat
            shutil.rmtree(directory, ignore_errors=True)
        self.web_urls[key] = info.file_id
        self.web_urls[web.url_key(page.url)] = info.file_id
        return info

    def remove(self, file_id: str) -> None:
        info = self.files.pop(file_id, None)
        if info:
            self.web_urls = {url: key for url, key in self.web_urls.items() if key != file_id}
            self.store.delete_file(self.session_id, file_id)
            for t in info.tables:
                self.tables.con.execute(f'DROP TABLE IF EXISTS "{t}"')
                self.tables.tables.pop(t, None)

    def close(self) -> None:
        self.store.delete_session(self.session_id)
        self.tables.con.close()

    def describe(self) -> str:
        """Compact inventory given to the router."""
        lines = []
        for f in self.doc_files:
            unit = "slides" if f.name.lower().endswith(".pptx") else "pages"
            lines.append(f"- {f.name} (document{f', {f.pages} {unit}' if f.pages else ''})")
        for f in self.sheet_files:
            lines.append(f"- {f.name} (spreadsheet; tables: {', '.join(f.tables)})")
        return "\n".join(lines) or "(no files uploaded)"
