"""Builds a deterministic, multi-format test corpus with known facts on known pages.

Used by the unit tests and the eval harness: `uv run python -m eval.corpus eval/corpus`
"""

import random
import sys
from pathlib import Path

import docx
import pandas as pd
import pymupdf
from docx.enum.text import WD_BREAK

REGIONS = ["North", "South", "East", "West"]
PRODUCTS = {"Atlas Arm": 64000, "Scout AMR": 38500, "Gripper Kit": 4200}

PDF_PAGES = [
    ("""<h1>Acme Robotics Annual Report 2025</h1>
<h2>1 Letter from the CEO</h2>
<p>Dear shareholders, 2025 was a breakthrough year. Under CEO Maria Lindqvist, Acme Robotics grew revenue
18% to $412.5 million, driven by demand for warehouse automation. We expanded to 1,240 employees across
four regions and signed our largest ever contract with Nordfrakt Logistics.</p>
<p>Our mission remains unchanged: make physical work safer and more productive.</p>"""),
    ("""<h2>2 Financial Highlights</h2>
<p>Net income reached $38.2 million, up from $21.9 million in 2024. Operating margin improved every quarter.</p>"""),
    ("""<h2>3 Products</h2>
<h3>3.1 Atlas Arm</h3>
<p>The Atlas Arm is a warehouse picking arm launched in March 2025 at a list price of $64,000. It picks
up to 1,100 items per hour with a 99.7% accuracy rate.</p>
<h3>3.2 Scout AMR</h3>
<p>Scout is our autonomous mobile robot. We shipped 1,850 Scout units in 2025, primarily to third-party
logistics providers.</p>"""),
    ("""<h2>4 Risks</h2>
<p>We depend on a single supplier of servo motors, Kinetix GmbH in Stuttgart. A disruption at Kinetix
could delay shipments by up to two quarters. We are qualifying a second supplier in 2026.</p>
<h2>5 Outlook</h2>
<p>For fiscal 2026 we guide to revenue of $480 to $500 million and plan to open an Austin facility.</p>"""),
]
QUARTERS = [("Q1", "92.1", "11.2%"), ("Q2", "98.4", "12.0%"), ("Q3", "104.7", "13.1%"), ("Q4", "117.3", "14.5%")]


def _pdf(path: Path) -> None:
    doc = pymupdf.open()
    for i, html in enumerate(PDF_PAGES):
        page = doc.new_page()
        page.insert_htmlbox(pymupdf.Rect(60, 60, 540, 190 if i == 1 else 520), html)
        if i == 1:  # draw a ruled table so the layout engine detects it
            y = 200
            rows = [("Quarter", "Revenue ($M)", "Operating margin")] + QUARTERS
            for r in rows:
                for j, cell in enumerate(r):
                    box = pymupdf.Rect(60 + j * 150, y, 210 + j * 150, y + 22)
                    page.draw_rect(box)
                    page.insert_text((box.x0 + 5, box.y1 - 7), cell, fontsize=10)
                y += 22
        page.insert_text((280, 800), f"Page {i + 1}", fontsize=8)
    doc.save(path)


def _docx(path: Path) -> None:
    d = docx.Document()
    d.add_heading("Employee Handbook", 0)
    d.add_heading("1 Working Hours", 1)
    d.add_paragraph("Core hours are 10:00 to 15:00. Employees may work remotely up to 3 days per week "
                    "with manager approval.")
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_heading("2 Leave Policy", 1)
    d.add_paragraph("Full-time employees receive 25 days of paid annual leave. Parental leave is 16 weeks, fully paid.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Leave type", "Days per year"
    for k, v in [("Sick leave", "12"), ("Bereavement", "5"), ("Volunteering", "2")]:
        r = t.add_row().cells
        r[0].text, r[1].text = k, v
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_heading("3 Expenses", 1)
    d.add_paragraph("Meals are reimbursed up to $75 per day when travelling. Submit receipts in Expensify "
                    "within 30 days of the expense.")
    d.save(path)


MD = """# Scout AMR FAQ

## Battery
Scout runs for 10 hours on a full charge and recharges in 90 minutes on the standard dock.

## Payload
The maximum payload is 150 kg. Loads above 120 kg reduce top speed to 1.2 m/s.

## Warranty
Scout ships with a 2-year warranty covering parts and labour. Extended warranties are available.
"""

TXT = """Board meeting minutes — 14 November 2025

Attendees: Maria Lindqvist (CEO), Daniel Okafor (CFO), Priya Raman (CTO), board members.

Decision: the board approved opening an Austin, Texas facility in Q2 2026 with 120 new hires.
Decision: the dividend remains suspended for 2026 to fund expansion.

Action items:
- Daniel Okafor to present a currency hedging plan at the January meeting.
- Priya Raman to report on second-supplier qualification for servo motors.
"""


def _orders() -> pd.DataFrame:
    rng = random.Random(7)
    rows = []
    for i in range(1, 121):
        product = rng.choice(list(PRODUCTS))
        units = rng.randint(1, 12)
        price = PRODUCTS[product]
        rows.append({
            "Order ID": f"SO-{1000 + i}",
            "Date": pd.Timestamp("2025-01-01") + pd.Timedelta(days=rng.randint(0, 364)),
            "Region": rng.choice(REGIONS),
            "Product": product,
            "Units": units,
            "Unit Price": price,
            "Revenue": units * price,
        })
    return pd.DataFrame(rows)


TARGETS = pd.DataFrame({"Region": REGIONS, "Target Revenue": [5_000_000, 4_000_000, 4_500_000, 3_500_000]})


def _xlsx(path: Path) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        # Title row + blank row above the real header: a common real-world mess.
        pd.DataFrame([["Acme Robotics — 2025 Orders"], [None]]).to_excel(xw, sheet_name="Orders", header=False,
                                                                         index=False)
        _orders().to_excel(xw, sheet_name="Orders", startrow=2, index=False)
        TARGETS.to_excel(xw, sheet_name="Targets", index=False)


def _customers() -> pd.DataFrame:
    rng = random.Random(11)
    segments = ["Enterprise", "Mid-market", "SMB"]
    return pd.DataFrame([{
        "customer_id": f"C{200 + i}",
        "name": f"Customer {i}",
        "region": rng.choice(REGIONS),
        "segment": rng.choice(segments),
        "signup_date": f"2024-{rng.randint(1, 12):02d}-{rng.randint(1, 28):02d}",
        "churned": rng.random() < 0.2,
    } for i in range(1, 51)])


def table_facts() -> dict[str, str]:
    """Ground-truth answers for table questions, formatted as they might appear in an answer."""
    o = _orders()
    by_region = o.groupby("Region")["Revenue"].sum()
    top_region = by_region.idxmax()
    atlas_units = int(o.loc[o["Product"] == "Atlas Arm", "Units"].sum())
    c = _customers()
    west_vs_target = by_region["West"] - int(TARGETS.set_index("Region").loc["West", "Target Revenue"])
    return {
        "total_revenue": f"{int(o['Revenue'].sum()):,}",
        "top_region": top_region,
        "top_region_revenue": f"{int(by_region[top_region]):,}",
        "atlas_units": str(atlas_units),
        "order_count": str(len(o)),
        "churned_customers": str(int(c["churned"].sum())),
        "enterprise_customers": str(int((c["segment"] == "Enterprise").sum())),
        "west_vs_target": f"{abs(int(west_vs_target)):,}",
        "west_vs_target_sign": "above" if west_vs_target > 0 else "below",
    }


def build_corpus(out: Path) -> dict[str, Path]:
    out.mkdir(parents=True, exist_ok=True)
    paths = {
        "pdf": out / "acme_annual_report_2025.pdf",
        "docx": out / "employee_handbook.docx",
        "md": out / "scout_faq.md",
        "txt": out / "board_minutes.txt",
        "xlsx": out / "sales_2025.xlsx",
        "csv": out / "customers.csv",
    }
    _pdf(paths["pdf"])
    _docx(paths["docx"])
    paths["md"].write_text(MD)
    paths["txt"].write_text(TXT)
    _xlsx(paths["xlsx"])
    _customers().to_csv(paths["csv"], index=False)
    return paths


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "eval/corpus")
    for kind, p in build_corpus(target).items():
        print(f"{kind:5} {p}")
